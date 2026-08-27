from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


FIXED_DATE = dt.datetime(2026, 8, 27, 0, 0, 0, tzinfo=dt.timezone.utc)
BLUE = "2E74B5"
DEEP_BLUE = "1F4D78"
GREY = "5B6573"
LIGHT = "F4F6F9"
PALE_BLUE = "EAF2F8"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_table_header(row) -> None:
    repeat_table_header(row)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def add_toc(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-1" \\h \\z '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Update field to display table of contents."
    separate.append(placeholder)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def clean_markdown_links(text: str) -> str:
    return re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", text)


def add_inline(paragraph, text: str, base_size: float | None = None) -> None:
    text = clean_markdown_links(text)
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*]+?\*)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            if base_size:
                run.font.size = Pt(base_size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor.from_string(DEEP_BLUE)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        if base_size:
            run.font.size = Pt(base_size)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        if base_size:
            run.font.size = Pt(base_size)


def configure_page(section, landscape: bool = False) -> None:
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Cm(29.7)
        section.page_height = Cm(21.0)
        section.left_margin = Cm(1.45)
        section.right_margin = Cm(1.45)
        section.top_margin = Cm(1.45)
        section.bottom_margin = Cm(1.45)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(1.9)
        section.right_margin = Cm(1.9)
        section.top_margin = Cm(1.75)
        section.bottom_margin = Cm(1.65)
    section.header_distance = Cm(0.75)
    section.footer_distance = Cm(0.75)


def configure_headers(section) -> None:
    header = section.header
    header.is_linked_to_previous = True
    p = header.paragraphs[0]
    p.text = "PROJECT EDGE / UKACTIVE   •   POINT-IN-TIME ETF ROTATION"
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(GREY)
    footer = section.footer
    footer.is_linked_to_previous = True
    f = footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = f.add_run("E2_DEVELOPMENTAL  •  HISTORICAL CUTOFF 2026-08-21  •  ")
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GREY)
    add_page_field(f)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    h1 = styles["Heading 1"]
    h1.font.name = "Calibri"
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor.from_string(BLUE)
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(10)
    h1.paragraph_format.keep_with_next = True

    h2 = styles["Heading 2"]
    h2.font.name = "Calibri"
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor.from_string(BLUE)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True

    h3 = styles["Heading 3"]
    h3.font.name = "Calibri"
    h3.font.size = Pt(12)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor.from_string(DEEP_BLUE)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.194)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208

    if "Source Citation" not in styles:
        style = styles.add_style("Source Citation", WD_STYLE_TYPE.PARAGRAPH)
    source = styles["Source Citation"]
    source.font.name = "Calibri"
    source.font.size = Pt(8)
    source.font.italic = True
    source.font.color.rgb = RGBColor.from_string(GREY)
    source.paragraph_format.space_before = Pt(4)
    source.paragraph_format.space_after = Pt(5)
    source.paragraph_format.line_spacing = 1.0

    if "Figure Caption UKACTIVE" not in styles:
        style = styles.add_style("Figure Caption UKACTIVE", WD_STYLE_TYPE.PARAGRAPH)
    caption = styles["Figure Caption UKACTIVE"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.color.rgb = RGBColor.from_string(GREY)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(9)
    caption.paragraph_format.keep_with_next = False

    if "Code Block UKACTIVE" not in styles:
        style = styles.add_style("Code Block UKACTIVE", WD_STYLE_TYPE.PARAGRAPH)
    code = styles["Code Block UKACTIVE"]
    code.font.name = "Consolas"
    code.font.size = Pt(8)
    code.paragraph_format.left_indent = Cm(0.5)
    code.paragraph_format.right_indent = Cm(0.5)
    code.paragraph_format.space_after = Pt(2)
    code.paragraph_format.line_spacing = 1.0

    # Word regenerates the TOC from these built-in styles. Compact spacing
    # keeps the final entry with the rest of the contents instead of creating
    # an almost-empty third TOC page.
    for level, size in ((1, 8.6), (2, 8.2), (3, 7.8)):
        name = f"TOC {level}"
        if name not in styles:
            continue
        toc_style = styles[name]
        toc_style.font.name = "Calibri"
        toc_style.font.size = Pt(size)
        toc_style.paragraph_format.space_before = Pt(0)
        toc_style.paragraph_format.space_after = Pt(1)
        toc_style.paragraph_format.line_spacing = 1.0


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(28)
    p.add_run("PROJECT EDGE / UKACTIVE RESEARCH PROGRAMME").font.color.rgb = RGBColor.from_string(BLUE)
    p.runs[0].font.size = Pt(11)
    p.runs[0].bold = True

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(25)
    title.paragraph_format.space_after = Pt(12)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = title.add_run("Development and Freeze of a Point-in-Time\nMulti-Horizon ETF Rotation Strategy for a UK SIPP")
    r.font.name = "Calibri"
    r.font.size = Pt(25)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(DEEP_BLUE)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(26)
    s = subtitle.add_run("Evidence from a Causal Industry-and-Theme ETF Universe, 2017–2026")
    s.font.name = "Calibri"
    s.font.size = Pt(15)
    s.font.color.rgb = RGBColor.from_string(GREY)

    metadata = [
        ("Paper version", "1.0"),
        ("Generation date", "2026-08-27"),
        ("Historical cutoff", "2026-08-21"),
        ("Frozen model", "INDUSTRY_PLUS_THEME | M2_BASE | TOP7 | MONTHLY | CASH0_ALWAYS_INVESTED | EQUAL"),
        ("Evidence classification", "E2_DEVELOPMENTAL"),
        ("Repository source", "df30c9be430c4a746a6e3ec0d97690612badd4fa"),
        ("Live readiness", "AWAITING_CURRENT_HOLDINGS_INPUT; pilot_operationally_ready=false"),
    ]
    table = doc.add_table(rows=len(metadata), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Cm(4.2)
    table.columns[1].width = Cm(12.0)
    for row, (label, value) in zip(table.rows, metadata):
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], LIGHT)
        for cell in row.cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(9)
        row.cells[0].paragraphs[0].runs[0].bold = True

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.left_indent = Cm(0.45)
    p.paragraph_format.right_indent = Cm(0.45)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), PALE_BLUE)
    p_pr.append(shd)
    run = p.add_run(
        "NON-DEPLOYMENT STATEMENT  —  This paper documents a frozen developmental research model and a conditional pilot design. "
        "It is not an investment recommendation, does not create broker orders, and does not establish operational readiness."
    )
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(DEEP_BLUE)
    doc.add_page_break()

    toc_heading = doc.add_paragraph("Contents", style="Heading 1")
    toc_heading.paragraph_format.space_before = Pt(0)
    toc = doc.add_paragraph()
    add_toc(toc)
    doc.add_page_break()


def split_table_row(line: str) -> list[str]:
    """Split a Markdown table row without treating code-span pipes as cells."""

    content = line.strip()
    if content.startswith("|"):
        content = content[1:]
    if content.endswith("|"):
        content = content[:-1]

    cells: list[str] = []
    current: list[str] = []
    in_code = False
    escaped = False
    for character in content:
        if escaped:
            current.append(character)
            escaped = False
            continue
        if character == "\\":
            escaped = True
            current.append(character)
            continue
        if character == "`":
            in_code = not in_code
            current.append(character)
            continue
        if character == "|" and not in_code:
            cells.append("".join(current).strip())
            current = []
            continue
        current.append(character)
    cells.append("".join(current).strip())
    return cells


def is_separator_row(cells: list[str]) -> bool:
    return all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def add_table(doc: Document, rows: list[list[str]], landscape: bool, caption: str | None = None) -> None:
    if landscape:
        section = doc.add_section(WD_SECTION.NEW_PAGE)
        configure_page(section, landscape=True)
        configure_headers(section)
    if caption:
        add_paragraph_block(doc, caption)
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    repeat_table_header(table.rows[0])
    font_size = 7.8 if cols >= 10 else 8.2 if cols >= 7 else 8.6
    for r_idx, source_row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            set_cell_margins(cell, 70, 85, 70, 85)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_idx == 0:
                set_cell_shading(cell, LIGHT)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            value = source_row[c_idx] if c_idx < len(source_row) else ""
            add_inline(p, value, base_size=font_size)
            if r_idx == 0:
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(DEEP_BLUE)
    if landscape:
        section = doc.add_section(WD_SECTION.NEW_PAGE)
        configure_page(section, landscape=False)
        configure_headers(section)


def add_paragraph_block(doc: Document, text: str) -> None:
    if text.startswith("> "):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.right_indent = Cm(0.35)
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(8)
        p_pr = p._p.get_or_add_pPr()
        p_bdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), BLUE)
        p_bdr.append(left)
        p_pr.append(p_bdr)
        add_inline(p, text[2:])
        return
    source_like = text.startswith("*Internal source") or text.startswith("*Sample") or text.startswith("*Current-only") or text.startswith("*Common sample") or text.startswith("*Samples")
    figure_caption = text.startswith("**Figure ")
    table_title = text.startswith("**Table ")
    if source_like:
        p = doc.add_paragraph(style="Source Citation")
        add_inline(p, text.strip("*"), base_size=8)
    elif figure_caption:
        p = doc.add_paragraph(style="Figure Caption UKACTIVE")
        add_inline(p, text, base_size=9)
    elif table_title:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        add_inline(p, text)
    else:
        p = doc.add_paragraph()
        add_inline(p, text)


def parse_markdown(doc: Document, markdown_path: Path) -> None:
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    start = next(i for i, line in enumerate(lines) if line.strip() == "## Abstract")
    i = start
    in_code = False
    code_lines: list[str] = []
    in_equation = False
    equation_lines: list[str] = []
    pending_table_caption: str | None = None
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code:
                p = doc.add_paragraph(style="Code Block UKACTIVE")
                p.paragraph_format.keep_together = True
                p_pr = p._p.get_or_add_pPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), LIGHT)
                p_pr.append(shd)
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(7.6)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if stripped == "\\[":
            in_equation = True
            equation_lines = []
            i += 1
            continue
        if stripped == "\\]" and in_equation:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(" ".join(part.strip() for part in equation_lines))
            run.font.name = "Cambria Math"
            run.font.size = Pt(10.5)
            in_equation = False
            equation_lines = []
            i += 1
            continue
        if in_equation:
            equation_lines.append(line)
            i += 1
            continue
        if not stripped:
            i += 1
            continue
        if stripped.startswith("**Table "):
            next_index = i + 1
            while next_index < len(lines) and not lines[next_index].strip():
                next_index += 1
            if next_index < len(lines) and lines[next_index].strip().startswith("|"):
                pending_table_caption = stripped
                i += 1
                continue
        image_match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            image_path = (markdown_path.parent / image_match.group(2)).resolve()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(image_path), width=Inches(6.35))
            i += 1
            continue
        if stripped.startswith("|") and "|" in stripped[1:]:
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(split_table_row(lines[i]))
                i += 1
            if len(table_lines) >= 2 and is_separator_row(table_lines[1]):
                table_lines.pop(1)
            add_table(
                doc,
                table_lines,
                landscape=len(table_lines[0]) >= 8,
                caption=pending_table_caption,
            )
            pending_table_caption = None
            continue
        heading_match = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading_match:
            marks, title = heading_match.groups()
            level = min(len(marks) - 1, 3)
            if level == 1 and re.match(r"\d+\.\s", title) and title not in {"1. Introduction"}:
                pass
            p = doc.add_paragraph(style=f"Heading {level}")
            add_inline(p, title)
            i += 1
            continue
        ordered = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if ordered:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, ordered.group(2))
            i += 1
            continue
        bullet = re.match(r"^-\s+(.+)$", stripped)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, bullet.group(1))
            i += 1
            continue
        add_paragraph_block(doc, stripped)
        i += 1


def normalize_docx(path: Path) -> None:
    temp = path.with_suffix(".normalized.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temp, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(source.namelist()):
            data = source.read(name)
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, data)
    temp.replace(path)


def build(markdown_path: Path, output_path: Path) -> None:
    doc = Document()
    configure_styles(doc)
    configure_page(doc.sections[0], landscape=False)
    configure_headers(doc.sections[0])
    add_cover(doc)
    parse_markdown(doc, markdown_path)

    core = doc.core_properties
    core.title = "Development and Freeze of a Point-in-Time Multi-Horizon ETF Rotation Strategy for a UK SIPP"
    core.subject = "Project EDGE / UKACTIVE — E2 developmental scientific paper"
    core.author = ""
    core.last_modified_by = ""
    core.keywords = "ETF rotation; industry momentum; point-in-time data; UK SIPP"
    core.comments = "Historical cutoff 2026-08-21. Frozen model. Non-deployment document."
    core.created = FIXED_DATE
    core.modified = FIXED_DATE
    core.last_printed = FIXED_DATE
    core.revision = 1

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    normalize_docx(output_path)
    digest = hashlib.sha256(output_path.read_bytes()).hexdigest()
    print(f"DOCX {output_path} sha256={digest}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--markdown", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.markdown.resolve(), args.output.resolve())


if __name__ == "__main__":
    main()
